"""Build the browser-side Office package templates used by transfusion.js.

The DOCX starts from the retained hospital form. The XLSX is authored by the
artifact-tool workflow and supplied as an input. Both packages are unpacked
into base64 entries so the extension can fill placeholders and rebuild valid
OOXML ZIP files without a remote service or a third-party runtime library.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = next((ROOT / "tests" / "fixtures").glob("*.docx"))
DEFAULT_WORK_DIR = ROOT / ".tmp-transfusion-export"
DEFAULT_OUTPUT = ROOT / "office-templates.js"


def set_font(run, *, size: float = 12, bold: bool = False, underline: bool = False) -> None:
    run.bold = bold
    run.underline = underline
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "Times New Roman")


def clear_runs(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_value_paragraph(paragraph, token: str, *, size: float = 10.5, center: bool = False) -> None:
    clear_runs(paragraph)
    run = paragraph.add_run(token)
    set_font(run, size=size, bold=True, underline=True)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_label_value(
    paragraph,
    label: str,
    token: str,
    *,
    label_size: float = 10.5,
    value_size: float = 10.5,
    newline: bool = False,
    center: bool = False,
) -> None:
    clear_runs(paragraph)
    label_run = paragraph.add_run(label)
    set_font(label_run, size=label_size)
    if newline:
        label_run.add_break()
    value_run = paragraph.add_run(token)
    set_font(value_run, size=value_size, bold=True, underline=True)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def get_cell(table, row: int, column: int):
    return table.rows[row].cells[column]


def build_docx_template(output_path: Path) -> None:
    document = Document(REFERENCE)
    table = document.tables[0]

    set_value_paragraph(get_cell(table, 1, 0).paragraphs[1], "{{patientFullName}}, {{patientBirthDate}} г.р.", size=9)
    set_value_paragraph(get_cell(table, 1, 3).paragraphs[1], "{{requestDate}} {{requestTime}}")
    set_value_paragraph(get_cell(table, 1, 7).paragraphs[1], "{{transfusionDate}}")
    set_value_paragraph(get_cell(table, 2, 0).paragraphs[1], "{{department}}")
    set_label_value(get_cell(table, 2, 3).paragraphs[0], "N и/б ", "{{patientHistoryNumber}}")
    set_label_value(get_cell(table, 2, 7).paragraphs[0], "Время начала трансфузии ", "{{startTime}}")
    set_label_value(get_cell(table, 2, 7).paragraphs[1], "Время окончания трансфузии ", "{{endTime}}")

    set_label_value(get_cell(table, 4, 0).paragraphs[0], "Группа крови реципиента AB0: ", "{{recipientAbo}}")
    set_label_value(get_cell(table, 4, 7).paragraphs[0], "Резус-принадлежность: ", "{{recipientRh}}")
    set_label_value(get_cell(table, 5, 0).paragraphs[0], "Антигены C, c, E, e, K: ", "{{recipientAntigens}}")
    set_label_value(get_cell(table, 5, 7).paragraphs[0], "Аллоиммунные антитела: ", "{{recipientAntibodies}}")
    set_value_paragraph(get_cell(table, 7, 0).paragraphs[0], "{{indication}}", center=True)

    set_label_value(get_cell(table, 9, 0).paragraphs[0], "Трансфузии компонентов крови в анамнезе: ", "{{previousTransfusions}}", label_size=9.5, value_size=9.5, center=True)
    set_label_value(get_cell(table, 9, 1).paragraphs[0], "Реакции и осложнения на трансфузии в анамнезе: ", "{{previousReactions}}", label_size=9.5, value_size=9.5, center=True)
    set_label_value(get_cell(table, 9, 9).paragraphs[0], "Трансфузии по индивидуальному подбору: ", "{{previousIndividualSelection}}", label_size=9.5, value_size=9.5, center=True)

    set_label_value(get_cell(table, 11, 0).paragraphs[0], "Наименование компонента донорской крови: ", "{{componentName}} {{componentDetails}}", label_size=9.5, value_size=9.5)
    set_label_value(get_cell(table, 11, 6).paragraphs[0], "Организация заготовки: ", "{{donorOrganization}}", label_size=9.5, value_size=9.5)
    set_label_value(get_cell(table, 12, 0).paragraphs[0], "Группа крови донора AB0: ", "{{donorAbo}} Rh: {{donorRh}}")
    set_label_value(get_cell(table, 12, 6).paragraphs[0], "Антигены эритроцитов донора C, c, E, e, K: ", "{{donorAntigens}}", label_size=9.5, value_size=9.5)
    set_value_paragraph(get_cell(table, 13, 0).paragraphs[1], "{{componentUnitNumber}}")
    set_value_paragraph(get_cell(table, 13, 1).paragraphs[1], "{{componentVolume}}")
    set_label_value(get_cell(table, 13, 6).paragraphs[0], "Дополнительные сведения: ", "{{componentNotes}}", label_size=8.5, value_size=9)
    set_label_value(get_cell(table, 14, 0).paragraphs[0], "Дата заготовки: ", "{{collectionDate}}")
    set_label_value(get_cell(table, 14, 6).paragraphs[0], "Срок годности: ", "{{expirationDate}}")

    set_label_value(get_cell(table, 16, 0).paragraphs[0], "Наименование медицинской организации, осуществившей индивидуальный подбор: ", "{{selectionOrganization}}", label_size=9, value_size=9)
    set_label_value(get_cell(table, 17, 0).paragraphs[0], "Дата исследования: ", "{{selectionDate}}")
    set_value_paragraph(get_cell(table, 18, 0).paragraphs[1], "{{selectionResponsible}}")
    set_value_paragraph(get_cell(table, 18, 4).paragraphs[1], "{{selectionConclusion}}")

    set_label_value(get_cell(table, 20, 0).paragraphs[0], "Наименования реагентов: ", "{{reagentName}}")
    set_label_value(get_cell(table, 21, 0).paragraphs[0], "N серии реагента: ", "{{reagentSeries}}")
    set_label_value(get_cell(table, 21, 10).paragraphs[0], "Срок годности: ", "{{reagentExpiration}}", label_size=10)
    set_value_paragraph(get_cell(table, 22, 0).paragraphs[1], "{{planeTest}}", center=True)
    set_value_paragraph(get_cell(table, 22, 5).paragraphs[1], "{{biologicalTest}}", center=True)
    set_label_value(get_cell(table, 24, 0).paragraphs[0], "Основные симптомы: ", "{{reactionSymptoms}}")
    set_label_value(get_cell(table, 24, 5).paragraphs[0], "Степень тяжести: ", "{{reactionSeverity}}")

    observation_tokens = [
        (27, 2, "{{beforeBp}}"), (27, 5, "{{beforePulse}}"),
        (27, 8, "{{beforeTemperature}}"), (27, 11, "{{beforeDiuresis}}"),
        (28, 2, "{{afterOneHourBp}}"), (28, 5, "{{afterOneHourPulse}}"),
        (28, 8, "{{afterOneHourTemperature}}"), (28, 11, "{{afterOneHourDiuresis}}"),
        (29, 2, "{{afterTwoHoursBp}}"), (29, 5, "{{afterTwoHoursPulse}}"),
        (29, 8, "{{afterTwoHoursTemperature}}"), (29, 11, "{{afterTwoHoursDiuresis}}"),
    ]
    for row, column, token in observation_tokens:
        set_value_paragraph(get_cell(table, row, column).paragraphs[0], token, size=10)

    set_label_value(get_cell(table, 30, 0).paragraphs[0], "Врач, осуществивший трансфузию: ", "{{transfusionDoctor}}")

    # Source example values are stored in Word content controls and are not
    # exposed through python-docx paragraph.runs. Every control in this form is
    # an example input slot, so remove any one that remains after slot rewrite.
    for content_control in list(table._tbl.findall(".//" + qn("w:sdt"))):
        content_control.getparent().remove(content_control)

    # The source contains a mandatory trailing paragraph after the full-page
    # table. Collapse it to one hidden point so Word keeps the form on one page.
    trailing = document.paragraphs[0]
    clear_runs(trailing)
    hidden_run = trailing.add_run(" ")
    set_font(hidden_run, size=1)
    vanish = OxmlElement("w:vanish")
    hidden_run._element.get_or_add_rPr().append(vanish)
    trailing.paragraph_format.space_before = Pt(0)
    trailing.paragraph_format.space_after = Pt(0)
    trailing.paragraph_format.line_spacing = Pt(1)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def package_entries(path: Path) -> list[dict[str, str]]:
    with zipfile.ZipFile(path) as archive:
        return [
            {
                "name": item.filename,
                "data": base64.b64encode(archive.read(item.filename)).decode("ascii"),
            }
            for item in archive.infolist()
            if not item.is_dir()
        ]


def write_browser_templates(docx_path: Path, xlsx_path: Path, output_path: Path) -> None:
    payload = {
        "referenceSha256": hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
        "docx": package_entries(docx_path),
        "xlsx": package_entries(xlsx_path),
    }
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    output_path.write_text(
        "(function publishFillBARSOfficeTemplates(globalScope) {\n"
        "    'use strict';\n"
        f"    globalScope.__FillBARSOfficeTemplates__ = {serialized};\n"
        "})(globalThis);\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--work-dir", type=Path, default=DEFAULT_WORK_DIR)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_WORK_DIR / "transfusion-template.xlsx")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    docx_path = args.work_dir / "transfusion-template.docx"
    build_docx_template(docx_path)
    write_browser_templates(docx_path, args.xlsx, args.output)
    print(f"DOCX template: {docx_path}")
    print(f"XLSX template: {args.xlsx}")
    print(f"Browser package: {args.output}")


if __name__ == "__main__":
    main()
